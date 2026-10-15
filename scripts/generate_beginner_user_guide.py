#!/usr/bin/env python3
"""Generate Orbis Equity beginner user guide (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("ORBIS EQUITY TERMINAL\n")
    tr.bold = True
    tr.font.size = Pt(28)
    tr.font.color.rgb = RGBColor(0x1A, 0x56, 0x8E)
    tr2 = title.add_run("Beginner Trader User Guide")
    tr2.bold = True
    tr2.font.size = Pt(20)
    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Equity Swing Terminal — concepts, logic, and step-by-step instructions\n"
        f"Document version: {date.today().isoformat()}"
    )
    doc.add_page_break()

    add_heading(doc, "Important disclaimer", 1)
    add_para(
        doc,
        "Orbis Equity (also referred to in project docs as KovaView) is a research and "
        "education terminal. It shows end-of-day (EOD) prices, rule-based screens, and "
        "paper portfolios. It does not connect to a broker and does not place real trades "
        "for you. Nothing in this guide or in the app is personal financial advice. "
        "Markets involve risk of loss. If you trade with real money, use a licensed broker "
        "and only risk capital you can afford to lose.",
    )
    doc.add_paragraph()

    add_heading(doc, "Who this guide is for", 1)
    add_para(
        doc,
        "This guide assumes you are new to trading vocabulary and to systematic equity "
        "tools. You will learn what the terminal measures, why those measures exist, and "
        "how to use each screen without getting lost in jargon.",
    )

    add_heading(doc, "Part 1 — Trading basics (in plain language)", 1)

    add_heading(doc, "What is a stock?", 2)
    add_para(
        doc,
        "A stock (share) is a small ownership slice of a public company. When you buy a "
        "share, you hope the business becomes more valuable over time, or that other "
        "traders will pay more for it soon. When you sell, you exchange that ownership "
        "for cash.",
    )

    add_heading(doc, "Long vs short", 2)
    add_bullets(
        doc,
        [
            "Long: You buy first and sell later. You profit if the price goes up.",
            "Short (paper only in this app): You sell borrowed shares first and buy them "
            "back later. You profit if the price falls. Real shorting requires a broker "
            "and borrow availability; Orbis only simulates shorts in some modules.",
        ],
    )

    add_heading(doc, "Price, volume, and the trading day", 2)
    add_bullets(
        doc,
        [
            "Open: First traded price when the exchange opens.",
            "High / Low: Highest and lowest prices during the session.",
            "Close: Last price when the regular session ends — most daily charts use this.",
            "Volume: How many shares traded. High volume often means more conviction.",
            "Percent change (day %): How much the close moved vs yesterday's close.",
        ],
    )

    add_heading(doc, "Swing trading vs day trading", 2)
    add_para(
        doc,
        "Orbis is built for swing trading: holding positions for several days to weeks "
        "based on daily (EOD) data. Day traders focus on intraday moves; only a few "
        "modules (Opening Range) use shorter bars. Most screens update after the US "
        "market close, not tick-by-tick.",
    )

    add_heading(doc, "Part 2 — How Orbis Equity works", 1)

    add_heading(doc, "What the terminal is", 2)
    add_para(
        doc,
        "Orbis Equity is a systematic equity research terminal. Each night, a data "
        "pipeline refreshes prices, fundamentals, and computed signals. The web app "
        "reads that database and presents modules: watchlist scanner, factor scores, "
        "sector rotation, paper portfolios, charts, and more.",
    )

    add_heading(doc, "End-of-day (EOD) data — read this first", 2)
    add_para(
        doc,
        "The banner at the top shows PRICES AS OF … EOD and EOD/DELAYED — NOT REAL-TIME. "
        "That means numbers reflect the last completed trading day, not live quotes. "
        "Open the STATUS page to see pipeline health and whether signals are suppressed "
        "when data is stale.",
    )

    add_heading(doc, "Market posture (equity weather)", 2)
    add_para(
        doc,
        "Posture is a 0–100 gauge summarizing how many stocks in the universe closed "
        "higher that day (advancers). Rough guide:",
    )
    add_bullets(
        doc,
        [
            "Above ~55: more stocks up than down — friendlier backdrop for long ideas.",
            "Around 50: mixed tape.",
            "Below ~45: more stocks down — defensive or selective stance.",
        ],
    )
    add_para(
        doc,
        "Some strategies (for example the Loop paper portfolio) skip new entries when "
        "posture is below 50. Posture is a context filter, not a guarantee.",
    )

    add_heading(doc, "The watchlist screener logic (Module 1)", 2)
    add_para(doc, "The watchlist finds large-cap names already moving with unusual volume:")
    add_bullets(
        doc,
        [
            "On list: day change ≥ +4% AND volume ≥ 2× the 50-day average volume.",
            "Hot: day change ≥ +10% AND volume ≥ 5× that average.",
            "Ranking: names that closed nearer the session high rank higher.",
            "Down day: negative daily change.",
        ],
    )
    add_para(
        doc,
        "This module does not place orders. It answers: “What is strong right now?” "
        "Execution ideas live in Loop, ORB, Bias, Cash, or your own broker.",
    )

    add_heading(doc, "Orbis Score (O score)", 2)
    add_para(
        doc,
        "On the screener and ticker pages you may see an Orbis Score from 1–99. It blends "
        "60% watchlist quality rank (momentum/near-high) with 40% fundamental composite "
        "score. Higher scores mean stronger alignment on both tape and fundamentals — "
        "still not a buy command.",
    )

    add_heading(doc, "Factor scores and F-Score (Module 2)", 2)
    add_para(doc, "Quant Fundamentals ranks each stock vs peers on five pillars:")
    add_bullets(
        doc,
        [
            "Value (25%): cheaper on metrics like P/E, earnings yield.",
            "Quality (25%): profitability, margins, Piotroski F-Score.",
            "Growth (20%): revenue and earnings growth.",
            "Earnings quality (15%): cash flow vs accounting earnings.",
            "Leverage (15%): debt and liquidity safety.",
        ],
    )
    add_para(
        doc,
        "Each factor is a percentile 0–100 within the universe. Composite is the weighted "
        "blend. Piotroski F-Score is 0–9 from nine yes/no financial health tests "
        "(profitability, leverage, efficiency). Scores of 7+ are often considered strong.",
    )

    add_heading(doc, "Part 3 — Navigating the app", 1)

    add_heading(doc, "Home cockpit", 2)
    add_bullets(
        doc,
        [
            "Equity Weather: posture gauge and pillar breakdown.",
            "Today's Brief: rule-based morning summary (best/worst sector).",
            "Market Pulse: advancers, on-list count, down-day share.",
            "Cross-Asset & Leadership: other markets and sector rotation snapshot.",
            "Show all module desks: links to every module with key metrics.",
        ],
    )

    add_heading(doc, "Top navigation bar", 2)
    add_para(doc, "Main links:")
    add_bullets(
        doc,
        [
            "WATCHLIST — momentum screener.",
            "EMA200 — names vs 200-day moving average trend.",
            "LOOP — Donchian breakout paper portfolio.",
            "ORB — opening range breakout (intraday-style, paper).",
            "SKEW — implied volatility map from option chains.",
            "ANALYZE — technical + fundamental vote.",
            "WEEKLY — weekly research view.",
            "BIAS — chart, levels, paper trade ideas.",
            "CASH — TEMA + Carver cash-sized sleeves.",
            "ROTATE — sector rotation and macro regime.",
            "QUANT — Quantropy risk and allocation math.",
            "FUNDAMENTALS — full factor screener.",
            "HEALTH — Health Care sector desk.",
            "EARNINGS — calendar and news sentiment.",
            "STATUS — data freshness and pipeline.",
        ],
    )

    add_heading(doc, "Jump to any ticker", 2)
    add_numbered(
        doc,
        [
            "Type a symbol (e.g. AAPL) in the Ticker box and press GO.",
            "Or press ⌘K (Mac) / Ctrl+K (Windows) to open the command palette.",
            "Click a symbol in any table to open its ticker detail page.",
        ],
    )

    add_heading(doc, "Focus universe (your personal list)", 2)
    add_para(
        doc,
        "You can narrow many modules to a custom list of symbols stored in your browser. "
        "Use this when you only care about a handful of names instead of the full "
        "S&P 500 / Nasdaq-100 universe.",
    )

    add_heading(doc, "Part 4 — Module guide (what each desk does)", 1)

    modules = [
        (
            "Watchlist (WATCHLIST)",
            "Scan for strong up days with volume. Filters: On list, Hot, All, Down day; "
            "region US/UK/EU; minimum day %. Use for idea generation only.",
        ),
        (
            "Quant Fundamentals (FUNDAMENTALS)",
            "Sort and filter by composite and factor scores, F-Score, sector. Good for "
            "finding quality or value names that are not necessarily moving today.",
        ),
        (
            "Loop Terminal (LOOP)",
            "Paper long-only portfolio: enter when close breaks above 55-day Donchian "
            "high and 20-day VWAP; exit via ATR chandelier stop. Rebalances every 5 "
            "sessions; max 8 names, sector caps. No new entries if posture < 50.",
        ),
        (
            "Earnings & News (EARNINGS)",
            "Upcoming earnings, recent beats/misses, and GDELT news headlines. Use to "
            "avoid surprise events or to read sentiment around a name.",
        ),
        (
            "Skew Map (SKEW)",
            "Implied volatility from listed options: skew, term structure, weekend vol. "
            "High IV often means the market expects big moves. Data is indicative, not "
            "a precise pricing model.",
        ),
        (
            "Opening Range (ORB)",
            "Paper long strategy on large gaps (≥4% vs prior close). Opening range is "
            "first 15 minutes NY time; buy break above range high with 1R bracket. "
            "Intraday logic — still driven by delayed/historical bars in the pipeline.",
        ),
        (
            "Stock Analysis (ANALYZE)",
            "Six technical signals (moving averages, MACD, RSI, etc.) vote BUY/SELL/HOLD "
            "with fundamental composite and F-Score shown alongside.",
        ),
        (
            "Quantropy (QUANT)",
            "Portfolio analytics: volatility, Sharpe, drawdown, CAPM beta, Altman Z, "
            "Markowitz-style allocations. Marked in-sample / diagnostic — not live trading.",
        ),
        (
            "Daily Bias (BIAS)",
            "Candlestick chart, pivot levels, ATR range, and paper ideas (fade levels, "
            "breakouts) from rules — not from a live order book.",
        ),
        (
            "TEMA + Carver Cash (CASH)",
            "Two paper sleeves on listed stocks at last close: TEMA 9/99/199 with MACD "
            "filter, and Carver EWMAC trend with vol targeting. Fully funded share sizing "
            "on a $100k paper book.",
        ),
        (
            "Beta Rotation (ROTATE)",
            "Macro canaries vote risk-on/off; sectors get labels (LEAD, LAG, etc.) from "
            "breadth, momentum, impulse; Carver D-rungs suggest how much to allocate. "
            "Diagnostic for sector trends, not a single-stock picker.",
        ),
        (
            "Health Sector (HEALTH)",
            "Health Care names with fundamentals and watchlist state — same ideas as "
            "Fundamentals but filtered to one sector.",
        ),
        (
            "Pharma (if enabled)",
            "Clinical trial pipeline signals for biotech/pharma names. Event-driven and "
            "high risk — treat as research, not tips.",
        ),
    ]
    for title, body in modules:
        add_heading(doc, title, 2)
        add_para(doc, body)

    add_heading(doc, "Ticker detail page", 2)
    add_para(
        doc,
        "Opening a symbol shows price chart, key metrics, factor breakdown, F-Score "
        "components, module-specific flags, and links to BIAS chart. Use this as your "
        "single-name dossier before doing anything in a real account.",
    )

    add_heading(doc, "Part 5 — Suggested workflows for beginners", 1)

    add_heading(doc, "Workflow A: Learn the market backdrop (5 minutes)", 2)
    add_numbered(
        doc,
        [
            "Open Home and read Equity Weather (posture) and Today's Brief.",
            "Check Market Pulse: advancers vs down day.",
            "Glance at Leadership & Rotation for leading sectors.",
            "Open STATUS if anything looks old or signals are off.",
        ],
    )

    add_heading(doc, "Workflow B: Find a stock to research (15 minutes)", 2)
    add_numbered(
        doc,
        [
            "Go to WATCHLIST with filter On list (or FUNDAMENTALS with Composite ≥ 70).",
            "Pick 2–3 symbols you recognize or in sectors you understand.",
            "Open each ticker page; note Orbis Score, day %, and F-Score.",
            "Open ANALYZE and BIAS for the same symbol — do technical and fundamental agree?",
            "Read EARNINGS for upcoming dates before any real trade.",
        ],
    )

    add_heading(doc, "Workflow C: Understand a systematic strategy (paper)", 2)
    add_numbered(
        doc,
        [
            "Read LOOP or CASH module descriptions in this guide.",
            "Open that module in the app and compare holdings to the rules.",
            "Track over weeks whether rules match your risk tolerance.",
            "Only then consider mirroring ideas manually in a broker — smaller size first.",
        ],
    )

    add_heading(doc, "Part 6 — Glossary", 1)
    glossary = [
        ("ATR", "Average True Range — typical daily price range; used for stops."),
        ("Advancers", "Stocks that closed up on the day."),
        ("Breadth", "How widespread gains or losses are across the universe."),
        ("CAPM / Beta", "How much a stock moves vs the market."),
        ("Donchian channel", "Highest high / lowest low over N days — breakout levels."),
        ("EOD", "End of day — after the regular session close."),
        ("EWMAC", "Exponentially weighted moving average crossover — trend signal."),
        ("GICS sector", "Standard industry grouping (e.g. Technology, Health Care)."),
        ("IV / Skew", "Implied volatility; skew compares put vs call demand."),
        ("MACD", "Momentum indicator from moving averages."),
        ("Paper portfolio", "Simulated positions — no real money."),
        ("Relative volume", "Today's volume divided by recent average."),
        ("RSI", "Relative Strength Index — overbought/oversold oscillator."),
        ("Sharpe ratio", "Return per unit of risk (higher is better, historically)."),
        ("VWAP", "Volume-weighted average price — intraday or rolling average."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Term"
    hdr[1].text = "Meaning"
    for term, meaning in glossary:
        row = table.add_row().cells
        row[0].text = term
        row[1].text = meaning

    doc.add_paragraph()
    add_heading(doc, "Part 7 — Frequently asked questions", 1)
    faq = [
        (
            "Why don't prices move when I watch during the day?",
            "Most data is EOD. The terminal is designed for swing decisions after the close.",
        ),
        (
            "Why do BUY/SELL counts show zero?",
            "Stale data or suppressed signals — check STATUS. The pipeline runs nightly (~19:30 UTC).",
        ),
        (
            "Can I connect my broker?",
            "No. Orbis is research-only. Place orders with your broker separately.",
        ),
        (
            "Which module should I trust most?",
            "None alone. Combine posture, fundamentals, and your own thesis. Disagreement "
            "between modules is a signal to slow down.",
        ),
    ]
    for q, a in faq:
        add_heading(doc, q, 3)
        add_para(doc, a)

    add_heading(doc, "Getting help", 1)
    add_para(
        doc,
        "For technical issues: STATUS page and repository DEPLOY.md. For learning: "
        "re-read Part 1 and the glossary, then paper-track one module for two weeks "
        "before risking capital.",
    )

    return doc


def main() -> None:
    out_paths = [
        Path("/workspace/docs/Orbis_Equity_Beginner_User_Guide.docx"),
        Path("/opt/cursor/artifacts/Orbis_Equity_Beginner_User_Guide.docx"),
    ]
    out_paths[0].parent.mkdir(parents=True, exist_ok=True)
    out_paths[1].parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    for p in out_paths:
        doc.save(str(p))
        print(f"Wrote {p}")


if __name__ == "__main__":
    main()
